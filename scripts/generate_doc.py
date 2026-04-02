#!/usr/bin/env python3
"""
Generate filled DOCX from Template.docx with subject details extracted from PDF.
Usage: python3 generate_doc.py <subject_code> <dept_code> <output_path>
Example: python3 generate_doc.py 20104-2001 20104 /tmp/output.docx
"""

import json
import os
import re
import sys
import tempfile
import urllib.request
import warnings

warnings.filterwarnings("ignore")

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.dirname(SCRIPT_DIR)
TEMPLATE_PATH = os.path.join(PROJECT_DIR, "Template.docx")
DEPT_FILE = os.path.join(PROJECT_DIR, "data", "departments.json")
SUBJECTS_FILE = os.path.join(PROJECT_DIR, "data", "subjects.json")

# Load data
with open(DEPT_FILE, "r", encoding="utf-8") as f:
    dept_data = json.load(f)
PDF_BASE_URL = dept_data["pdfBaseUrl"]

with open(SUBJECTS_FILE, "r", encoding="utf-8") as f:
    subj_data = json.load(f)


def find_department(dept_code):
    for d in dept_data["departments"]:
        if d["code"] == dept_code:
            return d
    return None


def find_subject(subject_code, dept_code):
    subjects = subj_data["subjects"].get(dept_code, [])
    for s in subjects:
        if s["code"] == subject_code:
            return s
    return None


# Thai character encoding fix (same as extract script)
def fix_thai_encoding(text):
    result = list(text)
    for i, ch in enumerate(result):
        prev_thai = i > 0 and '\u0e00' <= result[i-1] <= '\u0e7f'
        next_thai = i + 1 < len(result) and '\u0e00' <= result[i+1] <= '\u0e7f'
        next_ch = result[i+1] if i + 1 < len(result) else ''

        if prev_thai and ch in ('&', 'P') and (next_thai or i == len(result)-1):
            result[i] = '\u0e49'  # ้
        elif prev_thai and ch == ')' and (next_thai or i == len(result)-1):
            result[i] = '\u0e4c'  # ์
        elif prev_thai and ch == '"' and next_thai:
            result[i] = '\u0e48'  # ่
        elif prev_thai and ch == 'Q':
            if next_ch in 'ง\u0e0d':
                result[i] = '\u0e31'  # ั
            elif next_thai or i == len(result)-1:
                result[i] = '\u0e49'  # ้
        elif prev_thai and ch == '@' and (next_thai or i == len(result)-1):
            result[i] = '\u0e49'  # ้
        elif prev_thai and ch == "'" and next_thai:
            result[i] = '\u0e48'  # ่
        elif prev_thai and ch == '#' and next_thai:
            result[i] = '\u0e48'  # ่
        elif prev_thai and ch == '*' and (next_thai or i == len(result)-1):
            result[i] = '\u0e4c'  # ์
        elif prev_thai and ch == '0' and next_thai:
            # Context: '0' between Thai chars might be ้
            result[i] = '\u0e49'
        elif prev_thai and ch == '2' and next_thai:
            result[i] = '\u0e49'  # ้
        elif ch == 'q' and prev_thai and next_thai:
            result[i] = '\u0e31'  # ั  or could be ็
        elif ch == 't' and prev_thai and next_thai:
            result[i] = '\u0e47'  # ็
        elif ch == 'v' and prev_thai and next_thai:
            result[i] = '\u0e48'  # ่
    return ''.join(result)


def extract_subject_details(pdf_path, subject_code):
    """Extract detailed subject info from PDF."""
    import pdfplumber

    details = {
        "standardRef": "-",
        "learningOutcomes": "",
        "objectives": "",
        "competencies": "",
        "description": "",
        "pdfPage": 0  # แผ่นกระดาษ (1-based PDF page)
    }

    code_pattern = subject_code.replace("-", "[-–]")

    with pdfplumber.open(pdf_path) as pdf:
        # Find the page containing this subject's details
        found = False
        collecting = ""
        full_text = ""

        for page in pdf.pages:
            text = page.extract_text()
            if not text:
                continue

            # Check if this page contains our subject code with details
            if re.search(code_pattern, text):
                # Check if it's the detail section (has อ้างอิงมาตรฐาน or จุดประสงค์)
                if any(kw in text for kw in ['อ.างอิงมาตรฐาน', 'อ0างอิงมาตรฐาน', 'จุดประสงค', 'สมรรถนะรายวิชา', 'คําอธิบายรายวิชา', 'คำอธิบายรายวิชา']):
                    found = True
                    full_text = text
                    # page.page_number is 1-based = แผ่นกระดาษใน PDF viewer
                    details["pdfPage"] = page.page_number
                    # Also get next page (description might overflow)
                    page_idx = page.page_number  # 1-based, so index = page_number
                    if page_idx < len(pdf.pages):
                        next_text = pdf.pages[page_idx].extract_text()
                        if next_text and not re.search(r'\d{5}[-–]\d{4}', next_text[:50]):
                            full_text += "\n" + next_text
                    break

        if not found:
            return details

        # Apply Thai encoding fix
        full_text = fix_thai_encoding(full_text)

        # Parse sections
        lines = full_text.split('\n')

        sections = {
            'standardRef': [],
            'learningOutcomes': [],
            'objectives': [],
            'competencies': [],
            'description': []
        }

        current_section = None
        skip_first_line = False

        for line in lines:
            line_stripped = line.strip()
            if not line_stripped:
                continue

            # Detect section headers
            if re.search(r'อ.างอิงมาตรฐาน', line_stripped):
                current_section = 'standardRef'
                continue
            elif re.search(r'ผลลัพธ.การเรียนรู.ระดับรายวิชา', line_stripped):
                current_section = 'learningOutcomes'
                continue
            elif re.search(r'จุดประสงค.รายวิชา', line_stripped):
                current_section = 'objectives'
                # Skip "เพื่อให้" part
                continue
            elif re.search(r'สมรรถนะรายวิชา', line_stripped):
                current_section = 'competencies'
                continue
            elif re.search(r'คํ?าอธิบายรายวิชา', line_stripped):
                current_section = 'description'
                continue

            # Detect when we've reached the next subject (new code pattern)
            if current_section and re.match(r'\d{5}[-–]\d{4}', line_stripped):
                break

            # Skip header/footer lines
            if re.match(r'^\d+$', line_stripped):  # Page numbers
                continue
            if 'หลักสูตรประกาศนียบัตร' in line_stripped:
                continue
            if 'สาขาวิชา' in line_stripped and len(line_stripped) > 30:
                continue

            if current_section and line_stripped:
                sections[current_section].append(line_stripped)

        # Join sections
        details['standardRef'] = '\n'.join(sections['standardRef']).strip() or '-'
        details['learningOutcomes'] = '\n'.join(sections['learningOutcomes']).strip()
        details['objectives'] = '\n'.join(sections['objectives']).strip()
        details['competencies'] = '\n'.join(sections['competencies']).strip()
        details['description'] = '\n'.join(sections['description']).strip()

    return details


def fill_template(template_path, output_path, data):
    """Fill Template.docx with subject data, handling split runs."""
    from docx import Document

    doc = Document(template_path)

    for paragraph in doc.paragraphs:
        replace_placeholders_in_paragraph(paragraph, data)

    doc.save(output_path)
    return output_path


def replace_placeholders_in_paragraph(paragraph, data):
    """Replace {placeholder} in a paragraph, even when split across runs."""
    runs = paragraph.runs
    if not runs:
        return

    # Build full text from all runs
    full_text = ''.join(r.text for r in runs)

    # Check if any placeholder exists
    has_placeholder = False
    for key in data:
        if '{' + key + '}' in full_text:
            has_placeholder = True
            break

    if not has_placeholder:
        return

    # Replace all placeholders
    new_text = full_text
    for key, value in data.items():
        placeholder = '{' + key + '}'
        new_text = new_text.replace(placeholder, str(value))

    # Strategy: find run boundaries and map new text back
    # Simple approach: keep first run's formatting, clear others
    if len(runs) == 1:
        runs[0].text = new_text
    else:
        # Preserve first run's formatting
        runs[0].text = new_text
        for run in runs[1:]:
            run.text = ""


def main():
    if len(sys.argv) < 4:
        print("Usage: python3 generate_doc.py <subject_code> <dept_code> <output_path>", file=sys.stderr)
        sys.exit(1)

    subject_code = sys.argv[1]
    dept_code = sys.argv[2]
    output_path = sys.argv[3]

    # Find department and subject info
    dept = find_department(dept_code)
    if not dept:
        print(json.dumps({"error": f"Department {dept_code} not found"}))
        sys.exit(1)

    subj = find_subject(subject_code, dept_code)
    if not subj:
        print(json.dumps({"error": f"Subject {subject_code} not found in dept {dept_code}"}))
        sys.exit(1)

    # Parse credit (ท-ป-น)
    credit_parts = subj.get("credit", "0-0-0").replace("–", "-").split("-")
    theory = credit_parts[0] if len(credit_parts) > 0 else "0"
    practice = credit_parts[1] if len(credit_parts) > 1 else "0"
    credits = credit_parts[2] if len(credit_parts) > 2 else "0"

    # Determine program level text
    if dept["level"] == "ปวช.":
        program_level = "หลักสูตรประกาศนียบัตรวิชาชีพ พุทธศักราช 2567"
    else:
        program_level = "หลักสูตรประกาศนียบัตรวิชาชีพชั้นสูง พุทธศักราช 2567"

    # Download PDF and extract subject details
    pdf_url = PDF_BASE_URL + dept["pdf"]
    print(f"Downloading PDF: {dept['pdf']}...", file=sys.stderr)

    tmp_pdf = tempfile.NamedTemporaryFile(suffix='.pdf', delete=False)
    try:
        urllib.request.urlretrieve(pdf_url, tmp_pdf.name)
        print(f"Extracting details for {subject_code}...", file=sys.stderr)
        details = extract_subject_details(tmp_pdf.name, subject_code)
    except Exception as e:
        print(f"PDF download/parse error: {e}", file=sys.stderr)
        details = {
            "standardRef": "-",
            "learningOutcomes": "",
            "objectives": "",
            "competencies": "",
            "description": ""
        }
    finally:
        os.unlink(tmp_pdf.name)

    # Prepare template data
    template_data = {
        "programLevel": program_level,
        "vocationType": dept["category"],
        "occupationGroup": dept["group"],
        "department": dept["name"],
        "courseCode": subject_code,
        "courseName": subj["nameTh"],
        "theoryHours": theory,
        "practiceHours": practice,
        "credits": credits,
        "standardRef": details["standardRef"],
        "learningOutcomes": details["learningOutcomes"],
        "objectives": details["objectives"],
        "competencies": details["competencies"],
        "description": details["description"]
    }

    # Fill template
    print(f"Generating document...", file=sys.stderr)
    fill_template(TEMPLATE_PATH, output_path, template_data)

    # Output result as JSON
    result = {
        "success": True,
        "file": output_path,
        "pdfPage": details.get("pdfPage", 0),
        "subject": {
            "code": subject_code,
            "name": subj["nameTh"],
            "nameEn": subj.get("nameEn", ""),
            "dept": dept["name"],
            "level": dept["level"]
        }
    }
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
